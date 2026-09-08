# -*- coding: utf-8 -*-
"""
bonus_robot.py — автоматический расчёт бонусов по выполнению плана.

Вся бизнес-логика читается из книги-конфигурации ТП.xlsx (листы Спр_*).
Программа сама находит файлы отгрузок и плана в сетевых папках по маскам,
склеивает отгрузки, чистит факт, сопоставляет с планом (с иерархией
менеджер → супервайзер → директор) и формирует отдельный Excel-файл на
каждый город + сводный. Итоговые Факт/%/Бонус — живые формулы Excel.

Запуск:
    python bonus_robot.py                      # период берётся из Спр_Параметры
    python bonus_robot.py --period 2026-06     # явно указать период
    python bonus_robot.py --config "C:\\путь\\ТП.xlsx"

Зависимости: pandas, openpyxl   (pip install pandas openpyxl)
"""

import argparse
import calendar
import datetime as dt
import glob
import os
import re
import sys

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────────────────────────────────────
#  СЛУЖЕБНОЕ
# ─────────────────────────────────────────────────────────────────────────────

LOG_LINES = []

MANUAL_ROWS = 12
def log(msg, level="INFO"):
    line = f"[{dt.datetime.now():%H:%M:%S}] {level:5} | {msg}"
    print(line)
    LOG_LINES.append(line)


def die(msg):
    log(msg, "STOP")
    log("Расчёт остановлен.", "STOP")
    _flush_log()
    sys.exit(1)


def _flush_log():
    try:
        with open(os.path.join(OUT_DIR, f"_log_{PERIOD}.txt"), "w", encoding="utf-8") as f:
            f.write("\n".join(LOG_LINES))
    except Exception:
        pass


def yes(v):
    """Трактовка текстовых флагов ДА/НЕТ/1/True."""
    return str(v).strip().upper() in ("ДА", "YES", "TRUE", "1", "1.0", "+")


def num(v):
    """Безопасное приведение к числу (терпит пробелы, запятые, пустоты)."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("\xa0", "").replace(" ", "").replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(s) if s not in ("", "-", ".") else 0.0
    except ValueError:
        return 0.0


def norm_tokens(name):
    """Набор значимых токенов имени: ё→е, без префикса-маршрута, без отчества-инициалов."""
    if name is None or (isinstance(name, float) and pd.isna(name)):
        return set()
    s = str(name).lower().replace("ё", "е")
    s = s.split("(")[0]
    s = re.split(r"\bв\s*место\b|\bвместо\b", s)[0]
    s = re.sub(r"[^а-яa-z ]", " ", s)
    return {t for t in s.split() if len(t) > 1}


# ─────────────────────────────────────────────────────────────────────────────
#  ЧТЕНИЕ КОНФИГУРАЦИИ ИЗ ТП.xlsx
# ─────────────────────────────────────────────────────────────────────────────

def load_params(cfg):
    """Спр_Параметры → dict{параметр: значение}."""
    df = pd.read_excel(cfg, sheet_name="Спр_Параметры", header=0, dtype=str)
    df = df.dropna(subset=[df.columns[0]])
    return {str(k).strip(): (v if pd.notna(v) else "")
            for k, v in zip(df.iloc[:, 0], df.iloc[:, 1])}


def load_paths(cfg):
    """Спр_Пути → DataFrame источников."""
    df = pd.read_excel(cfg, sheet_name="Спр_Пути", header=0, dtype=str)
    df = df.rename(columns={
        df.columns[0]: "code", df.columns[1]: "kind",
        df.columns[2]: "folder", df.columns[3]: "mask", df.columns[4]: "active"})
    return df.dropna(subset=["code"])


def load_region_map(cfg):
    """Спр_Соответствия, блок 1: регион → (город, учитывать)."""
    raw = pd.read_excel(cfg, sheet_name="Спр_Соответствия", header=None, dtype=str)
    reg = {}
    start = None
    for i in range(len(raw)):
        if str(raw.iloc[i, 0]).strip() == "Регион в выгрузке":
            start = i + 1
            break
    if start is None:
        die("Спр_Соответствия: не найден заголовок «Регион в выгрузке».")
    for i in range(start, len(raw)):
        region = raw.iloc[i, 0]
        if pd.isna(region) or str(region).strip() == "":
            break
        city = raw.iloc[i, 1]
        count = raw.iloc[i, 3]
        reg[str(region).strip()] = {
            "city": None if pd.isna(city) else str(city).strip(),
            "count": yes(count)}
    return reg


def load_category_map(cfg):
    """Спр_Соответствия, блок 2: категория в источнике → канон. категория + роль."""
    raw = pd.read_excel(cfg, sheet_name="Спр_Соответствия", header=None, dtype=str)
    cat = {}
    start = None
    for i in range(len(raw)):
        if str(raw.iloc[i, 0]).strip() == "Категория в источнике":
            start = i + 1
            break
    if start is None:
        die("Спр_Соответствия: не найден заголовок «Категория в источнике».")
    for i in range(start, len(raw)):
        src = raw.iloc[i, 0]
        if pd.isna(src) or str(src).strip() == "":
            break
        canon = raw.iloc[i, 2]
        role = raw.iloc[i, 3]
        cat[str(src).strip().lower()] = {
            "canon": None if pd.isna(canon) else str(canon).strip(),
            "role": "" if pd.isna(role) else str(role).strip()}
    return cat


def _org_tokens(name):
    """Значимые слова наименования юрлица (без юр.форм): 'ТОО «Бытхим Трейд»' -> {'бытхим','трейд'}."""
    s = str(name).lower().replace("«", " ").replace("»", " ").replace('"', " ").replace("'", " ")
    s = re.sub(r"[^0-9a-zа-я ]", " ", s)
    stop = {"тоо", "ип", "пк", "ооо", "ао", "оао", "зао", "компания", "kz", "llp"}
    return frozenset(w for w in s.split() if len(w) > 2 and w not in stop)


def load_orgs(cfg):
    """Спр_Организации → внутригрупповые/тендерные коды + наборы слов наименований."""
    df = pd.read_excel(cfg, sheet_name="Спр_Организации", header=0, dtype=str)
    df = df.rename(columns={df.columns[0]: "code", df.columns[1]: "name"})
    internal, tender, name_sets = set(), set(), []
    for _, r in df.iterrows():
        internal_row = yes(r.get("Внутригрупповая"))
        code = r.get("code")
        if pd.notna(code) and str(code).strip():
            code = str(code).strip().lstrip("'")
            if internal_row:
                internal.add(code)
            if yes(r.get("Тендерная")):
                tender.add(code)
        if internal_row:
            ts = _org_tokens(r.get("name"))
            if ts:
                name_sets.append(ts)
    return internal, tender, name_sets


def load_rules(cfg):
    """Спр_Правила → список активных правил бонусов."""
    df = pd.read_excel(cfg, sheet_name="Спр_Правила", header=0)
    df = df.rename(columns={
        "Категория (канон.)": "canon", "Вид бонуса": "kind_bonus",
        "Тип правила": "rule", "Ставка бонуса": "rate",
        "База расчёта (ставка для расчёта плана)": "base",
        "Дедлайн": "deadline", "Шаг за 1% перевыпол.": "step",
        "Кэп %": "cap", "Лимит выплаты %": "limit",
        "Порог показателя": "thr", "Доля выплаты": "share", "Активно": "active"})
    df = df[df["active"].map(yes)]
    return df


def load_plan_bridge(cfg):
    """Спр_ПланМенеджеры → мост «ФИО в плане + город» → ID сотрудника (если заполнен)."""
    df = pd.read_excel(cfg, sheet_name="Спр_ПланМенеджеры", header=0, dtype=str)
    df = df.rename(columns={
        "Город": "city", "ФИО в файле планов": "fio",
        "ID сотрудника": "id", "Тип строки": "row_type",
        "Учитывать в расчёте": "use"})
    bridge = {}
    for _, r in df.iterrows():
        if pd.isna(r.get("fio")):
            continue
        key = (str(r.get("city")).strip(), frozenset(norm_tokens(r.get("fio"))))
        rid = r.get("id")
        if pd.notna(rid) and str(rid).strip():
            bridge[key] = str(rid).strip()
    return bridge


def load_canon_override(cfg):
    """Спр_ПланМенеджеры → (город, ФИО) → канон. категория (ручное переопределение)."""
    df = pd.read_excel(cfg, sheet_name="Спр_ПланМенеджеры", header=0, dtype=str)
    df = df.rename(columns={"Город": "city", "ФИО в файле планов": "fio",
                            "Категория (канон.)": "canon"})
    ov = {}
    for _, r in df.iterrows():
        if pd.isna(r.get("fio")) or pd.isna(r.get("canon")):
            continue
        canon = str(r["canon"]).strip()
        if canon in ("", "—", "— не сопоставлено"):
            continue
        ov[(str(r.get("city")).strip(), frozenset(norm_tokens(r.get("fio"))))] = canon
    return ov


# ─────────────────────────────────────────────────────────────────────────────
#  ПОИСК И ЧТЕНИЕ ФАЙЛОВ В СЕТЕВЫХ ПАПКАХ
# ─────────────────────────────────────────────────────────────────────────────

def period_range(period):
    """'2026-06' → ('01.06.2026-30.06.2026', год, месяц)."""
    y, m = int(period[:4]), int(period[5:7])
    last = calendar.monthrange(y, m)[1]
    return f"01.{m:02d}.{y}-{last:02d}.{m:02d}.{y}", y, m


def find_file(folder, mask, period_str):
    """Найти единственный файл в папке по маске с подстановкой {период}."""
    if folder is None or (isinstance(folder, float) and pd.isna(folder)) \
            or str(folder).strip().lower() in ("", "nan", "none"):
        return None, "путь к папке не указан"
    if not os.path.isdir(folder):
        return None, f"папка не найдена: {folder}"
    pat = str(mask).replace("{период}", period_str)
    hits = [f for f in glob.glob(os.path.join(folder, "*.xls*"))
            if _mask_match(os.path.basename(f), pat)
            and not os.path.basename(f).startswith("~")]
    if not hits:
        return None, f"нет файла по маске «{pat}» в {folder}"
    if len(hits) > 1:
        months = {"01": "январ", "02": "феврал", "03": "март", "04": "апрел",
                  "05": "май", "06": "июн", "07": "июл", "08": "август",
                  "09": "сентябр", "10": "октябр", "11": "ноябр", "12": "декабр"}
        mw = months.get(period_str[3:5], "")
        pref = [h for h in hits if mw and mw in os.path.basename(h).lower()]
        if pref:
            hits = pref
        else:
            log(f"в папке несколько файлов, месяц не распознан — беру первый: {folder}", "WARN")
    return sorted(hits)[0], None


def _mask_match(name, pattern):
    """Проверка имени по маске с * (регистронезависимо)."""
    rx = "^" + re.escape(pattern).replace(r"\*", ".*") + "$"
    return re.match(rx, name, flags=re.IGNORECASE) is not None


def read_shipments(path):
    """Чтение файла отгрузок: шапка во 2-й строке, коды — текст, фикс латинской 'c'."""
    df = pd.read_excel(path, sheet_name=0, header=1, dtype=str)
    df.columns = [str(c).replace("cотрудника", "сотрудника").strip() for c in df.columns]
    return df


def read_plan(path, category_map):
    """Файл плана «Распределение по торговой команде» → план по менеджерам."""
    df = pd.read_excel(path, sheet_name=0, header=0)
    df = df.iloc[:, :4]
    df.columns = ["city", "fio", "cat_src", "plan"]
    df["city"] = df["city"].ffill()
    df = df.dropna(subset=["fio"])
    df["plan"] = df["plan"].map(num)
    df["canon"] = df["cat_src"].map(
        lambda c: (category_map.get(str(c).strip().lower(), {}) or {}).get("canon"))
    df["tok"] = df["fio"].map(lambda x: frozenset(norm_tokens(x)))
    return df


# ─────────────────────────────────────────────────────────────────────────────
#  РАСЧЁТ
# ─────────────────────────────────────────────────────────────────────────────

def load_prihod_remap(cfg):
    """Спр_ПриходыКлиенты → перенос прихода клиента на другого менеджера/город."""
    try:
        df = pd.read_excel(cfg, sheet_name="Спр_ПриходыКлиенты", header=0, dtype=str)
    except Exception:
        return []
    df = df.rename(columns={df.columns[0]: "client", df.columns[1]: "manager", df.columns[2]: "city"})
    rules = []
    for _, r in df.iterrows():
        cl = r.get("client")
        if pd.isna(cl) or not str(cl).strip():
            continue
        rules.append({"key": str(cl).strip().lower(),
                      "manager": "" if pd.isna(r.get("manager")) else str(r.get("manager")).strip(),
                      "city": "" if pd.isna(r.get("city")) else str(r.get("city")).strip()})
    return rules


def load_magnum_rules(cfg):
    """Спр_ПриходыMagnum → партнёры, чей приход делится по филиалам пропорционально отгрузке."""
    try:
        df = pd.read_excel(cfg, sheet_name="Спр_ПриходыMagnum", header=0, dtype=str)
    except Exception:
        return []
    df = df.rename(columns={df.columns[0]: "key"})
    if len(df.columns) > 1:
        df = df.rename(columns={df.columns[1]: "active"})
    rules = []
    for _, r in df.iterrows():
        k = r.get("key")
        if pd.isna(k) or not str(k).strip():
            continue
        if "active" in df.columns and not yes(r.get("active")):
            continue
        words = frozenset(_name_words(k))
        if words:
            rules.append({"key": str(k).strip(), "words": words})
    return rules


def _name_words(s):
    """Слова наименования для сопоставления партнёра (регистр, ё, знаки — нормализованы)."""
    t = str(s or "").lower().replace("ё", "е")
    t = re.sub(r"[^0-9a-zа-я ]", " ", t)
    return {w for w in t.split() if w}

def read_prihod(path):
    """Файл приходов: 3 строки шапки. Приход = безнал+нал+эквайринг."""
    df = pd.read_excel(path, sheet_name=0, header=None, dtype=str)
    df = df.iloc[3:, :7].copy()
    df.columns = ["partner", "manager", "podr", "beznal", "nal", "acq", "dolg"]
    df = df.dropna(subset=["manager"])
    df["prihod"] = df["beznal"].map(num) + df["nal"].map(num) + df["acq"].map(num)
    return df


def aggregate_prihod(frames, region_map, remap=None, exclude_words=None):
    """Приходы по менеджеру: сумма (безнал+нал+эквайринг), город из Подразделения.
    remap — перенос прихода клиента на другого менеджера/город (Спр_ПриходыКлиенты).
    Сырое Подразделение сохраняется в колонке podr — нужно для листа ручного разбора."""
    if not frames:
        return pd.DataFrame(columns=["fio_prihod", "city", "podr", "tok", "prihod"])
    df = pd.concat(frames, ignore_index=True)
    remap = remap or []

     # партнёры с пропорциональным распределением исключаем из обычного матчинга по ФИО
    if exclude_words:
        pw = df["partner"].map(_name_words)
        keep = pw.map(lambda s: not any(kw <= s for kw in exclude_words))
        skipped = int((~keep).sum())
        if skipped:
            log(f"Приходы: {skipped} строк партнёров Magnum исключены из обычного матчинга "
                f"(распределяются пропорционально отгрузке)")
        df = df[keep]
        if len(df) == 0:
            return pd.DataFrame(columns=["fio_prihod", "city", "podr", "tok", "prihod"])

    def _eff(prow):
        ptoks = set(re.sub(r"[^0-9a-zа-я ]", " ", str(prow.get("partner", "")).lower()).split())
        man = str(prow.get("manager") or "")
        city = None
        for rule in remap:
            kw = set(rule["key"].split())
            if kw and kw <= ptoks:
                man = rule["manager"] or man
                city = rule["city"]
                break
        return man, city

    eff = [_eff(p) for _, p in df.iterrows()]
    df["_man_eff"] = [e[0] for e in eff]
    df["_city_forced"] = [e[1] for e in eff]
    skip = ("НЕАКТИВНЫЕ", "ПОСТАВЩИК", "ПОСТАВЩИКИ", "")
    rows = []
    for man, g in df.groupby(df["_man_eff"].fillna("").astype(str).str.strip()):
        mu = man.upper()
        if mu in skip or "СОТРУДНИК" in mu:
            continue
        podr_raw = str(g["podr"].dropna().iloc[0]).strip() if g["podr"].notna().any() else ""
        forced = [c for c in g["_city_forced"] if c]
        if forced:
            city = forced[0]
        else:
            city = (region_map.get(podr_raw, {}) or {}).get("city")
        rows.append({"fio_prihod": man, "city": city, "podr": podr_raw,
                     "tok": frozenset(norm_tokens(man)), "prihod": g["prihod"].sum()})
    return pd.DataFrame(rows)


def _match_prihod(row, prihod_agg):
    """Индекс строки приходов для менеджера: по ФИО (токены), город — как уточнение, не жёсткий фильтр."""
    if prihod_agg is None or len(prihod_agg) == 0:
        return None
    best_i, best = None, (0, 0)
    for i, pr in prihod_agg.iterrows():
        ov = len(row["tok"] & pr["tok"])
        if ov < 2:
            continue
        same_city = 1 if (pr["city"] and row["city"] and str(pr["city"]).strip() == str(row["city"]).strip()) else 0
        score = (ov, same_city)
        if score > best:
            best, best_i = score, i
    return best_i


def distribute_magnum(fact, prihod_raw, mag_rules, plan):
    """Приход партнёра делится по филиалам пропорционально ЕГО ЖЕ отгрузке.

    доля филиала = отгрузка филиала / отгрузка партнёра
    приход филиала = приход партнёра * доля
    Получатель — менеджер с наибольшей отгрузкой Magnum в этом городе,
    резерв — менеджер категории «ТП А» этого города.
    """
    if not mag_rules or fact is None or len(fact) == 0:
        return [], pd.DataFrame()
    cust_col = _find_col(fact, ["Контрагент"])
    man_col = _find_col(fact, ["Торговый менеджер"])
    if cust_col is None:
        log("Magnum: в отгрузках нет колонки «Контрагент» — распределение пропущено", "WARN")
        return [], pd.DataFrame()
    fw = fact[cust_col].map(_name_words)
    pw = prihod_raw["partner"].map(_name_words) if (prihod_raw is not None and len(prihod_raw)) else None

    add, detail = [], []
    for rule in mag_rules:
        kw = rule["words"]
        fsub = fact[fw.map(lambda s: kw <= s)]
        total_ship = fsub["_sum"].sum() if len(fsub) else 0.0
        if not total_ship:
            log(f"Magnum «{rule['key']}»: отгрузок за период не найдено — пропуск", "WARN")
            continue
        total_inc = float(prihod_raw.loc[pw.map(lambda s: kw <= s), "prihod"].sum()) if pw is not None else 0.0
        if not total_inc:
            log(f"Magnum «{rule['key']}»: приходов за период не найдено — распределять нечего", "WARN")
            continue
        log(f"Magnum «{rule['key']}»: отгрузка {total_ship:,.0f}, приход {total_inc:,.0f}")
        for city, g in fsub.groupby(fsub["_city"].fillna("")):
            city = str(city).strip()
            if not city:
                continue
            ship = g["_sum"].sum()
            share = ship / total_ship
            amount = total_inc * share
            fio, src = None, ""
            if man_col:
                by_man = g.groupby(g[man_col].fillna("").astype(str).str.strip())["_sum"].sum()
                by_man = by_man[by_man.index != ""].sort_values(ascending=False)
                if len(by_man):
                    fio, src = str(by_man.index[0]), "отгрузка Magnum"
            if not fio:
                cand = plan[(plan["city"].astype(str).str.strip() == city) &
                            (plan["canon"].astype(str).str.strip() == "ТП А")]
                if len(cand):
                    fio, src = str(cand.iloc[0]["fio"]), "ТП А города (резерв)"
            if fio:
                add.append({"city": city, "fio": fio, "prihod": amount})
            else:
                log(f"Magnum «{rule['key']}» / {city}: получатель не определён, "
                    f"{amount:,.0f} остались нераспределёнными", "WARN")
            detail.append({"Партнёр": rule["key"], "Город": city, "Отгрузка": ship,
                           "Доля": share, "Приход партнёра": total_inc,
                           "Приход города": amount,
                           "Отнесён на": fio or "— получатель не определён —",
                           "Основание": src or "нет"})
    return add, pd.DataFrame(detail)


def merge_magnum(prihod_agg, add):
    """Добавить распределённые суммы Magnum к приходу нужного менеджера."""
    if not add:
        return prihod_agg
    if prihod_agg is None:
        prihod_agg = pd.DataFrame(columns=["fio_prihod", "city", "podr", "tok", "prihod"])
    prihod_agg = prihod_agg.copy().reset_index(drop=True)
    for a in add:
        tok = frozenset(norm_tokens(a["fio"]))
        hit = None
        for i, pr in prihod_agg.iterrows():
            same_city = (not pr["city"] or not a["city"]
                         or str(pr["city"]).strip() == str(a["city"]).strip())
            if pr["tok"] == tok or (len(tok & pr["tok"]) >= 2 and same_city):
                hit = i
                break
        if hit is not None:
            prihod_agg.at[hit, "prihod"] = float(prihod_agg.at[hit, "prihod"]) + a["prihod"]
            prihod_agg.at[hit, "city"] = prihod_agg.at[hit, "city"] or a["city"]
            base = str(prihod_agg.at[hit, "podr"] or "")
            if "Magnum" not in base:
                prihod_agg.at[hit, "podr"] = (base + "; Magnum (распред.)").strip("; ")
        else:
            prihod_agg = pd.concat([prihod_agg, pd.DataFrame([{
                "fio_prihod": a["fio"], "city": a["city"], "podr": "Magnum (распред.)",
                "tok": tok, "prihod": a["prihod"]}])], ignore_index=True)
    return prihod_agg



def clean_fact(ship, region_map, internal, tender_codes, name_sets, params):
    """Разделить отгрузки на: факт менеджеров / исключено / тендер."""
    field = params.get("Поле суммы факта", "Сумма со скидкой")
    if field not in ship.columns:
        die(f"В отгрузках нет колонки суммы «{field}». Есть: {list(ship.columns)}")

    ship = ship.copy()
    ship["_sum"] = ship[field].map(num)
    if not yes(params.get("Учитывать возвраты (отрицательные суммы)", "ДА")):
        ship.loc[ship["_sum"] < 0, "_sum"] = 0.0

    code_col = _find_col(ship, ["Код контрагента"])
    reg_col = _find_col(ship, ["Регион"])
    man_col = _find_col(ship, ["Торговый менеджер"])
    cust_col = _find_col(ship, ["Контрагент"])
    ship["_code"] = ship[code_col].map(lambda x: str(x).strip().lstrip("'") if pd.notna(x) else "")

    _stop = {"тоо", "ип", "пк", "ооо", "ао", "оао", "зао", "компания", "kz", "llp"}

    def _cust_tokens(x):
        t = str(x).lower().replace("«", " ").replace("»", " ").replace('"', " ")
        t = re.sub(r"[^0-9a-zа-я ]", " ", t)
        return {w for w in t.split() if len(w) > 2 and w not in _stop}

    ship["_custtok"] = ship[cust_col].map(_cust_tokens) if cust_col else [set() for _ in range(len(ship))]

    tender_by_region = "тендер" in str(params.get("Признак тендера", "")).lower()

    def bucket(r):
        reg = str(r.get(reg_col, "")).strip()
        man = str(r.get(man_col, "")).strip().upper() if man_col else ""
        rm = region_map.get(reg)
        if man in ("ПОСТАВЩИКИ", "ПОСТАВЩИК"):
            return "excluded"
        if any(es <= r["_custtok"] for es in name_sets):
            return "excluded"
        if (tender_by_region and reg.lower().startswith("тендер")) or r["_code"] in tender_codes:
            return "tender"
        if r["_code"] in internal:
            return "excluded"
        if rm is None:
            return "excluded_unmapped"
        if rm.get("count") is False:
            return "excluded"
        return "fact"

    ship["_bucket"] = ship.apply(bucket, axis=1)
    ship["_city"] = ship[reg_col].map(lambda x: (region_map.get(str(x).strip(), {}) or {}).get("city"))

    unmapped = ship[ship["_bucket"] == "excluded_unmapped"]
    if len(unmapped):
        regs = sorted(unmapped[reg_col].dropna().astype(str).str.strip().unique())
        log(f"{len(unmapped)} строк с несопоставленным регионом исключены из факта. "
            f"Добавьте регион(ы) в Спр_Соответствия: {regs}", "WARN")
        ship.loc[ship["_bucket"] == "excluded_unmapped", "_bucket"] = "excluded"

    fact = ship[ship["_bucket"] == "fact"]
    excluded = ship[ship["_bucket"] == "excluded"]
    tender = ship[ship["_bucket"] == "tender"]
    return fact, excluded, tender


def _find_col(df, candidates):
    for c in candidates:
        if c in df.columns:
            return c
    for c in df.columns:
        if any(cand.lower() in str(c).lower() for cand in candidates):
            return c
    return None


def aggregate_fact(fact):
    """Факт по менеджеру: ключ = (ID, нормализованное имя, город)."""
    man_col = _find_col(fact, ["Торговый менеджер"])
    id_col = _find_col(fact, ["ID сотрудника"])
    rows = []
    grp = fact.groupby([fact[man_col].fillna(""), fact["_city"].fillna("")], dropna=False)
    for (man, city), g in grp:
        if str(man).strip() == "":
            continue
        ids = [str(x).strip() for x in g[id_col].dropna().unique() if str(x).strip()] if id_col else []
        rows.append({
            "fio_fact": man, "city": city or None,
            "id": ids[0] if ids else None,
            "tok": frozenset(norm_tokens(man)),
            "fact": g["_sum"].sum()})
    return pd.DataFrame(rows)


def _match_own(row, fact_agg, used):
    """Собственные отгрузки строки плана: по ID, затем по токенам имени + город."""
    if row.get("id_bridge"):
        cand = fact_agg[(fact_agg["id"] == row["id_bridge"]) & (~fact_agg.index.isin(used))]
        if len(cand):
            return cand.index[0]
    best_i, best_ov = None, 0
    for i, fr in fact_agg.iterrows():
        if i in used:
            continue
        if fr["city"] and row["city"] and str(fr["city"]).strip() != str(row["city"]).strip():
            continue
        ov = len(row["tok"] & fr["tok"])
        if ov > best_ov:
            best_ov, best_i = ov, i
    return best_i if best_ov >= 2 else None


def classify_row(cat_src, canon):
    """Роль строки плана: (role, group, scope). role=manager|sv|director; group=ВС|А; scope=личн|итого|None."""
    c = str(cat_src).strip().lower()
    if c.startswith("ноп"):
        return ("nop", "both", "личн" if "личн" in c else "итого")
    if "директор" in c or "дир." in c:
        return ("director", "both", "личн" if "личн" in c else "итого")
    if c.startswith("sv") or c.startswith("св") or "супервай" in c:
        grp = "А" if "кат" in c else "ВС"
        return ("sv", grp, "личн" if "личн" in c else "итого")
    cn = str(canon).strip()
    grp = "А" if (cn.startswith("ТП А") or "Магнум" in cn) else "ВС"
    return ("manager", grp, None)


def bonus_plan(pct, rule):
    """plan_proportional_cap: бонус за выполнение плана. pct = факт/план (доля)."""
    rate = num(rule["rate"])
    base = num(rule["base"])
    deadline = num(rule["deadline"])
    cap = num(rule["cap"])
    limit = num(rule["limit"])
    if pct < deadline:
        return 0.0
    if pct <= 1.0:
        return pct * rate
    if pct < cap:
        return rate + (pct - 1.0) * base
    return rate + base * limit


def build_results(plan, fact_agg, bridge, rules, prihod_agg, params):
    """Матчинг + иерархический роллап факта (менеджер→супервайзер→директор) + бонус."""
    round_to = 1 if "1" in str(params.get("Округление бонусов", "до 1 тенге")) else 0
    plan_rules = rules[rules["kind_bonus"].str.contains("Выполнение плана", na=False)]
    deb_rules = rules[rules["kind_bonus"].str.contains("Дебитор", na=False)]
    dir_rules = rules[rules["kind_bonus"].str.contains("Приходы", na=False)]
    spec_rules = rules[rules["kind_bonus"].str.contains("Специальн", na=False)]
    target_rules = rules[rules["kind_bonus"].str.contains("Целев", na=False)]

    plan = plan.copy().reset_index(drop=True)
    plan["id_bridge"] = plan.apply(
        lambda r: bridge.get((str(r["city"]).strip(), frozenset(r["tok"]))), axis=1)
    cls = plan.apply(lambda r: classify_row(r["cat_src"], r["canon"]), axis=1)
    plan["role"] = [x[0] for x in cls]
    plan["group"] = [x[1] for x in cls]
    plan["scope"] = [x[2] for x in cls]

    used, own, mname, used_prihod = set(), {}, {}, set()

    # 1) менеджеры забирают свои отгрузки (с расходованием строк факта)
    for i, r in plan[plan["role"] == "manager"].iterrows():
        fi = _match_own(r, fact_agg, used)
        if fi is not None:
            used.add(fi); own[i] = fact_agg.loc[fi, "fact"]; mname[i] = fact_agg.loc[fi, "fio_fact"]
        else:
            own[i] = 0.0; mname[i] = None
    # 2) собственные отгрузки супервайзеров, директоров и НОП (итоговые строки)
    for i, r in plan[(plan["role"].isin(["sv", "director", "nop"])) & (plan["scope"] != "личн")].iterrows():
        fi = _match_own(r, fact_agg, used)
        if fi is not None:
            used.add(fi); own[i] = fact_agg.loc[fi, "fact"]; mname[i] = fact_agg.loc[fi, "fio_fact"]
        else:
            own[i] = 0.0; mname[i] = None

    # 3) факт без плана — на контроль
    fact_only = fact_agg.loc[~fact_agg.index.isin(used)]
    fact_only = fact_only[fact_only["fact"] != 0]

    # 4) роллап по городам
    def team_fact(city, grp):
        idx = plan[(plan["role"] == "manager") & (plan["city"] == city) & (plan["group"] == grp)].index
        return sum(own.get(k, 0.0) for k in idx)

    def sv_own_sum(city, grp):
        idx = plan[(plan["role"] == "sv") & (plan["scope"] != "личн") &
                   (plan["city"] == city) & (plan["group"] == grp)].index
        return sum(own.get(k, 0.0) for k in idx)

    rows = []
    for i, r in plan.iterrows():
        if r["scope"] == "личн":       # личные строки SV/дир — только для роллапа, бонус не считаем
            continue
        city = r["city"]
        if r["role"] == "manager":
            fact = own.get(i, 0.0)
        elif r["role"] == "sv":
            fact = team_fact(city, r["group"]) + own.get(i, 0.0)
        elif r["role"] == "nop":
            fact = (team_fact(city, "ВС") + team_fact(city, "А")
                    + sv_own_sum(city, "ВС") + sv_own_sum(city, "А") + own.get(i, 0.0))
        else:  # director: обе команды + собств. отгрузки SV + собств. директора
            fact = (team_fact(city, "ВС") + team_fact(city, "А")
                    + sv_own_sum(city, "ВС") + sv_own_sum(city, "А") + own.get(i, 0.0))
        rec = {"city": city, "fio": r["fio"], "canon": r["canon"], "cat_src": r["cat_src"],
               "role": r["role"], "group": r["group"], "plan": r["plan"],
               "own": own.get(i, 0.0), "fact": fact,
               "fio_fact": mname.get(i), "id": r.get("id_bridge")}
        plan_val = num(r["plan"])
        if r["role"] == "nop" and not plan_val:      # план НОП = план директора города
            d = plan[(plan["role"] == "director") & (plan["city"] == city) &
                     (plan["scope"] != "личн")]
            if len(d):
                plan_val = num(d.iloc[0]["plan"])
        rec["plan"] = plan_val
        rec["pct"] = (fact / plan_val) if plan_val else 0.0
        rule = plan_rules[(plan_rules["Город"].astype(str).str.strip() == str(city).strip()) &
                          (plan_rules["canon"].astype(str).str.strip() == str(r["canon"]).strip())]
        if len(rule) and r["plan"]:
            rl = rule.iloc[0]
            rec["rate_plan"] = num(rl["rate"]); rec["base_plan"] = num(rl["base"])
            rec["deadline_"] = num(rl["deadline"]); rec["cap_"] = num(rl["cap"]); rec["limit_"] = num(rl["limit"])
            rec["bonus_plan"] = round(bonus_plan(rec["pct"], rl), round_to); rec["rule_ok"] = True
        else:
            rec["rate_plan"] = rec["base_plan"] = rec["bonus_plan"] = 0.0
            rec["deadline_"] = rec["cap_"] = rec["limit_"] = 0.0; rec["rule_ok"] = False
        drule = deb_rules[(deb_rules["Город"].astype(str).str.strip() == str(city).strip()) &
                          (deb_rules["canon"].astype(str).str.strip() == str(r["canon"]).strip())]
        pi = _match_prihod(r, prihod_agg)
        rec["prihod"] = float(prihod_agg.loc[pi, "prihod"]) if pi is not None else 0.0
        if pi is not None:
            used_prihod.add(pi)
        if len(drule):
            dl = drule.iloc[0]
            rec["deb_applies"] = True
            rec["deb_rate"] = num(dl["rate"]); rec["deb_thr"] = num(dl["thr"])
        else:
            rec["deb_applies"] = False
            rec["deb_rate"] = rec["deb_thr"] = 0.0
        if r["role"] == "director":
            drow = dir_rules[(dir_rules["Город"].astype(str).str.strip() == str(city).strip()) &
                             (dir_rules["canon"].astype(str).str.strip() == "Директор")]
            rec["dir_rate"] = num(drow.iloc[0]["share"]) if len(drow) else 0.0
        else:
            rec["dir_rate"] = 0.0
        srule = spec_rules[(spec_rules["Город"].astype(str).str.strip() == str(city).strip()) &
                           (spec_rules["canon"].astype(str).str.strip() == str(r["canon"]).strip())]
        rec["spec_rate"] = num(srule.iloc[0]["rate"]) if len(srule) else 0.0
        trule = target_rules[(target_rules["Город"].astype(str).str.strip() == str(city).strip()) &
                             (target_rules["canon"].astype(str).str.strip() == str(r["canon"]).strip())]
        rec["tgt_rate"] = num(trule.iloc[0]["rate"]) if len(trule) else 0.0
        rows.append(rec)
    if prihod_agg is not None and len(prihod_agg):
        prihod_only = prihod_agg.loc[~prihod_agg.index.isin(used_prihod)]
        prihod_only = prihod_only[prihod_only["prihod"] != 0]
    else:
        prihod_only = pd.DataFrame(columns=["fio_prihod", "podr", "city", "prihod"])
    return pd.DataFrame(rows), fact_only, prihod_only


# ─────────────────────────────────────────────────────────────────────────────
#  ФОРМИРОВАНИЕ ОТЧЁТОВ
# ─────────────────────────────────────────────────────────────────────────────

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Заголовки листа несопоставленных приходов
PR_COLS = {"fio_prihod": "Менеджер в приходах", "podr": "Подразделение",
           "city": "Город", "prihod": "Приход"}


def _style_sheet(ws, df, title, warn_pct):
    """Живой лист: Факт/%/Бонус — формулы Excel. Служебные колонки свёрнуты в группу."""
    cols = [
        ("city", "Город", 13), ("fio", "ФИО (план)", 32), ("canon", "Категория", 12),
        ("role", "Роль", 10), ("group", "Группа", 9), ("rate_plan", "Ставка", 11),
        ("base_plan", "База", 11), ("deadline_", "Дедлайн", 9), ("cap_", "Кэп", 8), ("limit_", "Лимит", 8),
        ("plan", "План", 14), ("own", "Собств. факт", 14), ("fact_f", "Факт (итого)", 14),
        ("pct_f", "% вып.", 9), ("bonus_f", "Бонус за план", 15),
        ("fio_fact", "ФИО в отгрузках", 26), ("rule_ok", "Правило", 8),
        ("prihod", "Приходы", 13), ("prihod_tot", "Приход итого", 13),
        ("deb_real", "Реализация", 13), ("deb_vz", "Взаимозачёты", 13),
        ("deb_pct", "% деб", 9), ("deb_rate", "Ставка деб", 12), ("deb_bonus", "Бонус деб", 14),
        ("dir_income", "Бонус от прихода", 15), ("dir_tender_in", "Приход тендер", 14),
        ("dir_tender_bonus", "Тендерный бонус", 15),
        ("spec_plan", "Спец задача План", 14), ("spec_fact", "Спец задача Факт", 14),
        ("spec_rate", "Спец задача Ставка", 14), ("spec_bonus", "Спец задача Бонус", 14),
        ("tgt_plan", "Целевая План", 14), ("tgt_fact", "Целевая Факт", 14),
        ("tgt_rate", "Целевая Ставка", 14), ("tgt_bonus", "Целевая Бонус", 14),
        ("total_bonus", "Итого бонус", 14),
    ]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    hr = 3
    for j, (_, name, w) in enumerate(cols, start=1):
        c = ws.cell(hr, j, name)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(hr + 1, 1)

    d0 = hr + 1
    d1 = hr + max(len(df), 1)
    OWN = f"$L${d0}:$L${d1}"; ROLE = f"$D${d0}:$D${d1}"; GRP = f"$E${d0}:$E${d1}"; CITY = f"$A${d0}:$A${d1}"

    r = d0
    for _, row in df.iterrows():
        ws.cell(r, 1, row.get("city"))
        ws.cell(r, 2, row.get("fio"))
        ws.cell(r, 3, row.get("canon"))
        role = str(row.get("role"))
        ws.cell(r, 4, role)
        ws.cell(r, 5, row.get("group"))
        ws.cell(r, 6, num(row.get("rate_plan")))
        ws.cell(r, 7, num(row.get("base_plan")))
        ws.cell(r, 8, num(row.get("deadline_")))
        ws.cell(r, 9, num(row.get("cap_")))
        ws.cell(r, 10, num(row.get("limit_")))
        ws.cell(r, 11, num(row.get("plan")))
        ws.cell(r, 12, num(row.get("own")))
        if role == "manager":
            f_fact = f"=L{r}"
        elif role == "sv":
            f_fact = f'=SUMIFS({OWN},{ROLE},"manager",{GRP},E{r},{CITY},A{r})+L{r}'
        else:
            f_fact = f'=SUMIFS({OWN},{ROLE},"manager",{CITY},A{r})+SUMIFS({OWN},{ROLE},"sv",{CITY},A{r})+L{r}'
        ws.cell(r, 13, f_fact)
        ws.cell(r, 14, f"=IF(K{r}=0,0,M{r}/K{r})")
        ws.cell(r, 15, f"=IF(N{r}<H{r},0,IF(N{r}<=1,N{r}*F{r},IF(N{r}<I{r},F{r}+(N{r}-1)*G{r},F{r}+G{r}*J{r})))")
        ws.cell(r, 16, row.get("fio_fact") or "")
        ws.cell(r, 17, "OK" if row.get("rule_ok") else "нет")
        Rown = f"$R${d0}:$R${d1}"
        ws.cell(r, 18, num(row.get("prihod")))
        ws.cell(r, 18).number_format = '# ##0;(# ##0);-'
        if role == "manager":
            ws.cell(r, 19, f"=R{r}")
        elif role == "sv":
            ws.cell(r, 19, f'=SUMIFS({Rown},{ROLE},"manager",{GRP},E{r},{CITY},A{r})+R{r}')
        else:
            ws.cell(r, 19, f'=SUMIFS({Rown},{ROLE},"manager",{CITY},A{r})+SUMIFS({Rown},{ROLE},"sv",{CITY},A{r})+R{r}')
        ws.cell(r, 19).number_format = '# ##0;(# ##0);-'
        if row.get("deb_applies"):
            Trng, Urng = f"$T${d0}:$T${d1}", f"$U${d0}:$U${d1}"
            if role == "sv" and str(row.get("group")) == "ВС":
                ws.cell(r, 22, f"=IFERROR(S{r}/M{r},0)")
            elif role == "sv":
                vz = f'SUMIFS({Urng},{ROLE},"manager",{GRP},E{r},{CITY},A{r})+U{r}'
                rl = f'SUMIFS({Trng},{ROLE},"manager",{GRP},E{r},{CITY},A{r})+T{r}'
                ws.cell(r, 22, f"=IFERROR((S{r}+{vz})/({rl}),0)")
            else:
                ws.cell(r, 22, f"=IFERROR((S{r}+U{r})/T{r},0)")
            ws.cell(r, 23, num(row.get("deb_rate")))
            ws.cell(r, 24, f"=IF(V{r}>{num(row.get('deb_thr'))},W{r},0)")
            ws.cell(r, 36, f"=O{r}+X{r}")
            ws.cell(r, 20).fill = PatternFill("solid", fgColor="FFF2CC")
            ws.cell(r, 21).fill = PatternFill("solid", fgColor="FFF2CC")
            for j in (20, 21, 23, 24, 36):
                ws.cell(r, j).number_format = '# ##0;(# ##0);-'
            ws.cell(r, 22).number_format = "0.0%"
        else:
            ws.cell(r, 36, f"=O{r}")
            ws.cell(r, 36).number_format = '# ##0;(# ##0);-'
        if role == "director":
            dr = num(row.get("dir_rate"))
            ws.cell(r, 25, f"=S{r}*{dr}")
            ws.cell(r, 25).number_format = '# ##0;(# ##0);-'
            if str(row.get("city")).strip() in ("Астана", "Шымкент", "Караганда"):
                ws.cell(r, 26).fill = PatternFill("solid", fgColor="FFF2CC")
                ws.cell(r, 26).number_format = '# ##0;(# ##0);-'
                ws.cell(r, 27, f"=Z{r}*0.01*(1-0.11)")
                ws.cell(r, 27).number_format = '# ##0;(# ##0);-'
            ws.cell(r, 36, f"=O{r}+Y{r}+AA{r}")
            ws.cell(r, 36).number_format = '# ##0;(# ##0);-'
            ws.cell(r, 30, num(row.get("spec_rate")))
            ws.cell(r, 30).number_format = '# ##0;(# ##0);-'
            ws.cell(r, 34, num(row.get("tgt_rate")))
            ws.cell(r, 34).number_format = '# ##0;(# ##0);-'
        for j in range(1, len(cols) + 1):
            ws.cell(r, j).border = BORDER
        for j in (6, 7, 11, 12, 13, 15):
            ws.cell(r, j).number_format = '# ##0;(# ##0);-'
        ws.cell(r, 14).number_format = "0.0%"
        for j in (8, 9, 10):
            ws.cell(r, j).number_format = "0.00"
        ws.cell(r, 12).fill = PatternFill("solid", fgColor="FFF2CC")
        if not row.get("rule_ok") or not row.get("fio_fact"):
            for j in range(1, len(cols) + 1):
                if j != 12:
                    ws.cell(r, j).fill = WARN_FILL
        r += 1

    ws.cell(r, 3, "ИТОГО").font = Font(bold=True)
    ws.cell(r, 12, f"=SUM(L{d0}:L{r-1})").number_format = '# ##0;(# ##0);-'
    ws.cell(r, 15, f"=SUM(O{d0}:O{r-1})").number_format = '# ##0;(# ##0);-'
    ws.cell(r, 18, f"=SUM(R{d0}:R{r-1})").number_format = '# ##0;(# ##0);-'
    ws.cell(r, 24, f"=SUM(X{d0}:X{r-1})").number_format = '# ##0;(# ##0);-'
    ws.cell(r, 25, f"=SUM(Y{d0}:Y{r-1})").number_format = '# ##0;(# ##0);-'
    ws.cell(r, 27, f"=SUM(AA{d0}:AA{r-1})").number_format = '# ##0;(# ##0);-'
    ws.cell(r, 36, f"=SUM(AJ{d0}:AJ{r-1})").number_format = '# ##0;(# ##0);-'
    for j in range(1, len(cols) + 1):
        ws.cell(r, j).fill = TOTAL_FILL
        ws.cell(r, j).border = BORDER
    ws.cell(r, 12).font = Font(bold=True); ws.cell(r, 15).font = Font(bold=True)

    for col in ("D", "E", "H", "I", "J", "P", "Q", "R",
                "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI"):
        ws.column_dimensions[col].outline_level = 1
        ws.column_dimensions[col].hidden = True
    ws.sheet_properties.outlinePr.summaryRight = False


def _dump_sheet(ws, df, title, cols):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12, color="C00000")
    if df is None or df.empty:
        ws["A3"] = "— нет строк —"
        return
    numeric = {"fact", "prihod", "prihod_tot", "_sum", "Приход"}
    show = [c for c in cols if c in df.columns]
    for j, name in enumerate(show, start=1):
        c = ws.cell(3, j, name)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        ws.column_dimensions[get_column_letter(j)].width = 24
    for i, (_, row) in enumerate(df.iterrows(), start=4):
        for j, name in enumerate(show, start=1):
            v = row.get(name, "")
            if name in numeric:
                cell = ws.cell(i, j, num(v))
                cell.number_format = '# ##0;(# ##0);-'
            else:
                ws.cell(i, j, "" if v is None else str(v))
                
def _manual_block(ws, row, n=12):
    """Блок ручного ввода внизу городского листа: мерчендайзеры и водители."""
    ws.cell(row, 2, "РУЧНОЙ ВВОД: мерчендайзеры и водители (автоматически не считается)")
    ws.cell(row, 2).font = Font(bold=True, color="C00000")
    hr = row + 1
    for j, nm in ((1, "Город"), (2, "ФИО"), (3, "Категория"), (15, "Бонус"), (36, "Итого бонус")):
        c = ws.cell(hr, j, nm)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
    for i in range(hr + 1, hr + 1 + n):
        for j in (1, 2, 3, 15):
            cc = ws.cell(i, j)
            cc.border = BORDER
            cc.fill = PatternFill("solid", fgColor="FFF2CC")
            cc.font = Font(color="0000FF")
        ws.cell(i, 15).number_format = '# ##0;(# ##0);-'
        ws.cell(i, 36, f"=O{i}").number_format = '# ##0;(# ##0);-'
        ws.cell(i, 36).border = BORDER
    t = hr + 1 + n
    ws.cell(t, 3, "ИТОГО ручной ввод").font = Font(bold=True)
    ws.cell(t, 15, f"=SUM(O{hr + 1}:O{t - 1})").number_format = '# ##0;(# ##0);-'
    ws.cell(t, 36, f"=SUM(AJ{hr + 1}:AJ{t - 1})").number_format = '# ##0;(# ##0);-'
    return t

def _city_from_podr(podr, region_map, cities):
    """Город для строки приходов: справочник → название без префикса «Филиал/Тендер» → вхождение."""
    s = str(podr or "").strip()
    if not s:
        return None
    m = region_map.get(s)
    if m and m.get("city"):
        return m["city"]
    t = re.sub(r"^\s*(филиал|тендер|подразделение|отдел)\s+", "", s, flags=re.IGNORECASE).strip()
    for c in cities:
        if t.lower() == str(c).strip().lower():
            return c
    for c in cities:
        cc = str(c).strip().lower()
        if len(cc) > 3 and cc in s.lower():
            return c
    return None


def _resolve_prihod_cities(prihod_only, region_map, cities):
    """Проставить городам приходов значения из cities; что не опознано — в сборный город."""
    if prihod_only is None or len(prihod_only) == 0:
        return prihod_only, None
    fb = [c for c in cities if "филиал" in str(c).lower()]
    fallback = fb[0] if fb else (cities[0] if cities else None)
    d = prihod_only.copy()
    if "podr" not in d.columns:
        d["podr"] = ""
    resolved, moved = [], 0
    for _, r in d.iterrows():
        c = r.get("city") if (r.get("city") and str(r.get("city")).strip() in map(str, cities)) else None
        if not c:
            c = _city_from_podr(r.get("podr"), region_map, cities)
        if not c or str(c).strip() not in map(str, cities):
            c = fallback
            moved += 1
        resolved.append(c)
    d["city"] = resolved
    if moved:
        log(f"{moved} строк приходов без опознанного города отнесены в файл «{fallback}»", "WARN")
    return d, fallback



def _prihod_sheet(wb, prihod_only, city=None):
    """Лист несопоставленных приходов.
    city=None → ВСЕ строки (в т.ч. с неопределённым городом, они иначе теряются).
    city='Алматы' → только этот город."""
    if prihod_only is None or len(prihod_only) == 0:
        d = prihod_only
    elif city is None:
        d = prihod_only.copy()
    else:
        d = prihod_only[prihod_only["city"].fillna("").astype(str).str.strip()
                        == str(city).strip()].copy()
    if d is not None and len(d):
        d["city"] = d["city"].fillna("")
        if "podr" not in d.columns:
            d["podr"] = ""
        d = d.sort_values("prihod", ascending=False).rename(columns=PR_COLS)
    ws = wb.create_sheet("Приходы не сопост.")
    _dump_sheet(ws, d, "Приходы менеджеров, которых НЕТ в плане — проверить и распределить вручную",
                list(PR_COLS.values()))
    if d is not None and len(d):
        r = 4 + len(d)
        ws.cell(r, 1, "ИТОГО").font = Font(bold=True)
        ws.cell(r, 4, f"=SUM(D4:D{r - 1})").number_format = '# ##0;(# ##0);-'
        ws.cell(r, 4).font = Font(bold=True)
        ws.cell(r, 4).fill = TOTAL_FILL
        ws.cell(r, 1).fill = TOTAL_FILL
    return ws


def _magnum_sheet(wb, magnum_detail, city=None):
    """Расшифровка распределения прихода Magnum по филиалам."""
    d = magnum_detail
    if d is not None and len(d) and city is not None:
        d = d[d["Город"].astype(str).str.strip() == str(city).strip()]
    ws = wb.create_sheet("Распределение Magnum")
    ws["A1"] = "Приход Magnum распределён по филиалам пропорционально отгрузке партнёра"
    ws["A1"].font = Font(bold=True, size=12, color="1F4E78")
    if d is None or len(d) == 0:
        ws["A3"] = "— нет строк —"
        return ws
    cols = ["Партнёр", "Город", "Отгрузка", "Доля", "Приход партнёра", "Приход города",
            "Отнесён на", "Основание"]
    widths = [24, 14, 16, 10, 18, 18, 30, 22]
    for j, (nm, w) in enumerate(zip(cols, widths), start=1):
        c = ws.cell(3, j, nm)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    r = 4
    for _, row in d.iterrows():
        ws.cell(r, 1, str(row["Партнёр"]))
        ws.cell(r, 2, str(row["Город"]))
        ws.cell(r, 3, num(row["Отгрузка"])).number_format = '# ##0;(# ##0);-'
        ws.cell(r, 4, num(row["Доля"])).number_format = "0.00%"
        ws.cell(r, 5, num(row["Приход партнёра"])).number_format = '# ##0;(# ##0);-'
        ws.cell(r, 6, num(row["Приход города"])).number_format = '# ##0;(# ##0);-'
        ws.cell(r, 7, str(row["Отнесён на"]))
        ws.cell(r, 8, str(row["Основание"]))
        for j in range(1, len(cols) + 1):
            ws.cell(r, j).border = BORDER
        r += 1
    ws.cell(r, 2, "ИТОГО").font = Font(bold=True)
    ws.cell(r, 3, f"=SUM(C4:C{r-1})").number_format = '# ##0;(# ##0);-'
    ws.cell(r, 6, f"=SUM(F4:F{r-1})").number_format = '# ##0;(# ##0);-'
    for j in range(1, len(cols) + 1):
        ws.cell(r, j).fill = TOTAL_FILL
        ws.cell(r, j).border = BORDER
    ws.cell(r, 3).font = Font(bold=True)
    ws.cell(r, 6).font = Font(bold=True)
    return ws


def write_city_report(city, dfc, excluded, tender, fact_only, prihod_only, period, warn_pct,
                      magnum_detail=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Бонусы"
    _style_sheet(ws, dfc, f"Бонусы за выполнение плана — {city} — {period}", warn_pct)
    _manual_block(ws, ws.max_row + 2, MANUAL_ROWS)

    nm = fact_only[(fact_only["city"] == city)] if fact_only is not None and len(fact_only) else fact_only
    _dump_sheet(wb.create_sheet("Не сопоставлено"), nm,
                "Есть факт, но не найден план — заполните ID в Спр_ПланМенеджеры",
                ["fio_fact", "city", "id", "fact"])
    _prihod_sheet(wb, prihod_only, city=city)
    _magnum_sheet(wb, magnum_detail, city=city)
    dump_cols = ["Торговый менеджер", "Регион", "Контрагент", "Код контрагента", "_sum"]
    _dump_sheet(wb.create_sheet("Исключено"),
                excluded[excluded["_city"].fillna("") == city] if len(excluded) else excluded,
                "Исключённые строки (внутригрупповые / межфирменные / не учитываемые регионы)", dump_cols)
    _dump_sheet(wb.create_sheet("Тендер"),
                tender[tender["_city"].fillna("") == city] if len(tender) else tender,
                "Тендерные отгрузки (в бонусы не входят)", dump_cols)

    path = os.path.join(OUT_DIR, f"Бонусы_{city}_{period}.xlsx")
    wb.save(path)
    log(f"Готов файл города: {os.path.basename(path)}  "
        f"(строк: {len(dfc)}, бонус: {dfc['bonus_plan'].sum():,.0f})")
    return path


def write_summary(all_df, period, prihod_only=None, magnum_detail=None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводный"
    _style_sheet(ws, all_df.sort_values(["city", "canon", "fio"]),
                 f"Сводный отчёт по бонусам — {period}", None)

    ws2 = wb.create_sheet("По городам")
    ws2["A1"] = "Итоги по городам"
    ws2["A1"].font = Font(bold=True, size=12, color="1F4E78")
    heads = ["Город", "Строк", "Бонус за план"]
    for j, h in enumerate(heads, 1):
        c = ws2.cell(3, j, h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        ws2.column_dimensions[get_column_letter(j)].width = 16
    r = 4
    for city, g in all_df.groupby("city"):
        ws2.cell(r, 1, city)
        ws2.cell(r, 2, len(g))
        ws2.cell(r, 3, g["bonus_plan"].sum()).number_format = '# ##0'
        r += 1

    # все несопоставленные приходы одним листом — независимо от города
    _prihod_sheet(wb, prihod_only, city=None)
    _magnum_sheet(wb, magnum_detail, city=None)

    path = os.path.join(OUT_DIR, f"Сводный_{period}.xlsx")
    wb.save(path)
    log(f"Готов сводный файл: {os.path.basename(path)}")
    return path

def write_curator_report(prihod_raw, region_map, period, params):
    """Отдельный файл: бонус куратора нескольких городов.

    База — ВЕСЬ приход этих городов из файла приходов (без привязки к менеджерам).
    Бонус от прихода = итого приход * 1%
    Тендерный бонус  = приход тендер * 1% - (приход тендер * 1% * 11%)   [тендер вводится вручную]
    """
    fio = str(params.get("Куратор: ФИО", "") or "").strip()
    cities_raw = str(params.get("Куратор: города", "") or "").strip()
    if not fio or not cities_raw:
        log("Куратор: в Спр_Параметры не заданы «Куратор: ФИО» / «Куратор: города» — файл не создаётся")
        return None
    cities = [c.strip() for c in re.split(r"[;,]", cities_raw) if c.strip()]

    per_city = {c: 0.0 for c in cities}
    tender_sum = 0.0
    if prihod_raw is not None and len(prihod_raw):
        for _, rr in prihod_raw.iterrows():
            podr = str(rr.get("podr") or "").strip()
            val = num(rr.get("prihod"))
            if re.match(r"^\s*тендер\b", podr, flags=re.IGNORECASE):
                tender_sum += val
                continue
            c = _city_from_podr(podr, region_map, cities)
            if c in per_city:
                per_city[c] += val
    if tender_sum:
        log(f"Куратор: тендерные подразделения ({tender_sum:,.0f}) в общий приход не включены — "
            f"тендер вносится вручную")

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Бонус куратора"
    ws["A1"] = f"Бонус куратора — {fio} — {period}"
    ws["A1"].font = Font(bold=True, size=13, color="1F4E78")
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 20

    for j, nm in enumerate(["Город", "Приход за период"], start=1):
        c = ws.cell(3, j, nm)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center")

    r = 4
    for city in cities:
        ws.cell(r, 1, city).border = BORDER
        cc = ws.cell(r, 2, per_city[city])
        cc.number_format = '# ##0;(# ##0);-'
        cc.border = BORDER
        r += 1
    first, last = 4, r - 1
    ws.cell(r, 1, "ИТОГО приход").font = Font(bold=True)
    t = ws.cell(r, 2, f"=SUM(B{first}:B{last})")
    t.number_format = '# ##0;(# ##0);-'
    t.font = Font(bold=True)
    for j in (1, 2):
        ws.cell(r, j).fill = TOTAL_FILL
        ws.cell(r, j).border = BORDER
    row_total = r

    r += 2
    ws.cell(r, 1, "Бонус от прихода (1%)").border = BORDER
    b1 = ws.cell(r, 2, f"=B{row_total}*0.01")
    b1.number_format = '# ##0;(# ##0);-'; b1.border = BORDER
    row_b1 = r
    r += 1
    ws.cell(r, 1, "Приход по тендеру (ручной ввод)").border = BORDER
    m = ws.cell(r, 2, 0)
    m.number_format = '# ##0;(# ##0);-'
    m.fill = PatternFill("solid", fgColor="FFF2CC")
    m.font = Font(color="0000FF")
    m.border = BORDER
    row_tender = r
    r += 1
    ws.cell(r, 1, "Тендерный бонус (1% минус 11%)").border = BORDER
    b2 = ws.cell(r, 2, f"=B{row_tender}*0.01-(B{row_tender}*0.01*0.11)")
    b2.number_format = '# ##0;(# ##0);-'; b2.border = BORDER
    row_b2 = r
    r += 1
    ws.cell(r, 1, "ИТОГО БОНУС").font = Font(bold=True)
    tot = ws.cell(r, 2, f"=B{row_b1}+B{row_b2}")
    tot.number_format = '# ##0;(# ##0);-'
    tot.font = Font(bold=True)
    for j in (1, 2):
        ws.cell(r, j).fill = TOTAL_FILL
        ws.cell(r, j).border = BORDER

    ws.cell(r + 2, 1, "Жёлтая ячейка — заполняется вручную. Остальное считается формулами.")
    ws.cell(r + 2, 1).font = Font(italic=True, size=9, color="808080")

    path = os.path.join(OUT_DIR, f"Бонус {fio}_{period}.xlsx")
    wb.save(path)
    log(f"Готов файл куратора: {os.path.basename(path)} "
        f"(города: {', '.join(cities)}; приход {sum(per_city.values()):,.0f})")
    return path


# ─────────────────────────────────────────────────────────────────────────────
#  ГЛАВНЫЙ КОНВЕЙЕР
# ─────────────────────────────────────────────────────────────────────────────

def main():
    global PERIOD, OUT_DIR

    ap = argparse.ArgumentParser(description="Робот расчёта бонусов")
    ap.add_argument("--config", default="ТП.xlsx", help="путь к книге-конфигурации ТП.xlsx")
    ap.add_argument("--period", default=None, help="период ГГГГ-ММ (по умолчанию из Спр_Параметры)")
    ap.add_argument("--out", default="Результат", help="папка для итоговых файлов")
    args = ap.parse_args()

    cfg = args.config
    if not os.path.isfile(cfg):
        print(f"Не найдена книга-конфигурация: {cfg}")
        sys.exit(1)

    params = load_params(cfg)
    global MANUAL_ROWS
    MANUAL_ROWS = int(num(params.get("Резерв строк: мерчендайзеры / водители", 12))) or 12
    PERIOD = args.period or str(params.get("Период расчёта", "")).strip()
    if not re.match(r"^\d{4}-\d{2}$", PERIOD):
        print(f"Некорректный период: '{PERIOD}'. Ожидается ГГГГ-ММ.")
        sys.exit(1)
    OUT_DIR = os.path.join(args.out, PERIOD)
    os.makedirs(OUT_DIR, exist_ok=True)

    period_str, _, _ = period_range(PERIOD)
    log(f"=== Расчёт бонусов за {PERIOD} ({period_str}) ===")

    paths = load_paths(cfg)
    region_map = load_region_map(cfg)
    category_map = load_category_map(cfg)
    internal, tender_codes, name_sets = load_orgs(cfg)
    rules = load_rules(cfg)
    bridge = load_plan_bridge(cfg)
    prihod_remap = load_prihod_remap(cfg)
    log(f"Конфигурация прочитана: правил активно {len(rules)}, "
        f"внутригрупп. кодов {len(internal)}, юрлиц по наименованию {len(name_sets)}, мостов ID {len(bridge)}.")

    # 1. Отгрузки из активных баз
    ship_sources = paths[(paths["kind"].str.strip().str.lower() == "отгрузки") &
                         (paths["active"].map(yes))]
    frames = []
    for _, s in ship_sources.iterrows():
        f, err = find_file(s["folder"], s["mask"], period_str)
        if err:
            die(f"Источник «{s['code']}»: {err}")
        log(f"Отгрузки «{s['code']}»: {os.path.basename(f)}")
        d = read_shipments(f)
        d["_source"] = s["code"]
        frames.append(d)
    if not frames:
        die("Не найдено ни одной активной базы отгрузок в Спр_Пути.")
    ship = pd.concat(frames, ignore_index=True)
    log(f"Склеено строк отгрузок: {len(ship)}")

    # 2. Чистка факта
    fact, excluded, tender = clean_fact(ship, region_map, internal, tender_codes, name_sets, params)
    log(f"Факт менеджеров: {len(fact)} строк | исключено: {len(excluded)} | тендер: {len(tender)}")

    # 3. Агрегация факта
    fact_agg = aggregate_fact(fact)

    # 3b. Приходы (дебиторка) — если включено в Спр_Параметры
    prihod_frames = []
    calc_deb = any("дебитор" in str(k).lower() and yes(v) for k, v in params.items())
    if calc_deb:
        pr_sources = paths[(paths["kind"].str.strip().str.lower() == "приходы") &
                           (paths["active"].map(yes))]
        for _, sp in pr_sources.iterrows():
            f, err = find_file(sp["folder"], sp["mask"], period_str)
            if err:
                log(f"Приходы «{sp['code']}»: {err} — пропущено", "WARN")
                continue
            log(f"Приходы «{sp['code']}»: {os.path.basename(f)}")
            prihod_frames.append(read_prihod(f))
    

    # 4. План
    plan_src = paths[(paths["kind"].str.strip().str.lower() == "план") & (paths["active"].map(yes))]
    if plan_src.empty:
        die("В Спр_Пути нет активного источника с типом «План».")
    ps = plan_src.iloc[0]
    pf, err = find_file(ps["folder"], ps["mask"], period_str)
    if err:
        die(f"Файл плана: {err}. Укажите папку плана в Спр_Пути (строка «План БМ»).")
    log(f"План: {os.path.basename(pf)}")
    plan = read_plan(pf, category_map)
    canon_ov = load_canon_override(cfg)
    if canon_ov:
        plan["canon"] = [canon_ov.get((str(c).strip(), frozenset(t)), k)
                         for c, t, k in zip(plan["city"], plan["tok"], plan["canon"])]
        log(f"Переопределений категорий из Спр_ПланМенеджеры: {len(canon_ov)}")

    # 4b. Приходы Magnum — делятся по филиалам пропорционально отгрузке партнёра
    mag_rules = load_magnum_rules(cfg)
    prihod_raw = pd.concat(prihod_frames, ignore_index=True) if prihod_frames else None
    magnum_add, magnum_detail = distribute_magnum(fact, prihod_raw, mag_rules, plan)
    prihod_agg = aggregate_prihod(prihod_frames, region_map, prihod_remap,
                                  exclude_words=[r["words"] for r in mag_rules])
    prihod_agg = merge_magnum(prihod_agg, magnum_add)
    if len(prihod_agg):
        log(f"Приходы собраны по {len(prihod_agg)} менеджерам, сумма {prihod_agg['prihod'].sum():,.0f}")

    # 5. Сопоставление, иерархический роллап факта и расчёт
    result, fact_only, prihod_only = build_results(plan, fact_agg, bridge, rules, prihod_agg, params)
    no_plan = len(fact_only)
    no_rule = (~result["rule_ok"]).sum()
    log(f"Рассчитано строк: {len(result)} | без плана (факт есть): {no_plan} "
        f"| без правила/канон.категории: {no_rule}")

    # 6. Отчёты по городам + сводный
    warn_pct = num(params.get("Порог предупреждения: % выполнения выше", 0)) or None
    cities = [c for c in result["city"].dropna().unique()]
    prihod_only, _fb = _resolve_prihod_cities(prihod_only, region_map, cities)
    if prihod_only is not None and len(prihod_only):
        log(f"Приходы без плана: {len(prihod_only)} менеджеров на сумму "
            f"{prihod_only['prihod'].sum():,.0f} — по листам «Приходы не сопост.»", "WARN")
    for city in cities:
        write_city_report(city, result[result["city"] == city],
                          excluded, tender, fact_only, prihod_only, PERIOD, warn_pct,
                          magnum_detail)
    write_summary(result, PERIOD, prihod_only, magnum_detail)
    write_curator_report(prihod_raw, region_map, PERIOD, params)


    log(f"ИТОГО бонус за план: {result['bonus_plan'].sum():,.0f} тенге по {len(cities)} городам.")
    log(f"Файлы сохранены в папку: {os.path.abspath(OUT_DIR)}")
    log("=== Готово ===")
    _flush_log()


if __name__ == "__main__":
    PERIOD, OUT_DIR = "", "."
    main()